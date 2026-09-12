"""
Automated unit & integration test for Document Extraction and Health Validation Service.
"""

import os
import io
import json
import asyncio
from app.services.document_extractor_service import (
    validate_document_health,
    extract_content,
    extract_legal_entities,
)

def run_tests():
    print("=== Step 1: Testing Health & Anti-Stub Validation ===")

    # Test 1A: Reject 0-byte empty stub
    h_empty = validate_document_health(b"", "empty_document.pdf")
    assert not h_empty["is_healthy"], "0-byte file should be rejected"
    assert h_empty["status"] == "REJECTED_EMPTY_STUB"
    print("  [OK] 0-byte stub correctly rejected:", h_empty["error"])

    # Test 1B: Reject unsupported format (.exe)
    h_exe = validate_document_health(b"MZ\x90\x00\x03\x00\x00\x00", "payload.exe")
    assert not h_exe["is_healthy"], ".exe should be rejected"
    assert h_exe["status"] == "REJECTED_UNSUPPORTED_FORMAT"
    print("  [OK] Unsupported extension correctly rejected:", h_exe["error"])

    # Test 1C: Reject corrupted fake PDF
    h_fake_pdf = validate_document_health(b"This is not a real PDF file header but has .pdf extension", "fake.pdf")
    assert not h_fake_pdf["is_healthy"]
    assert h_fake_pdf["status"] == "REJECTED_CORRUPTED_HEADER"
    print("  [OK] Fake PDF without magic bytes rejected:", h_fake_pdf["error"])

    print("\n=== Step 2: Testing DOCX Generation, Health & Extraction ===")
    import docx
    docx_doc = docx.Document()
    docx_doc.add_heading("IN THE 105TH DISTRICT COURT OF NUECES COUNTY, TEXAS", level=1)
    docx_doc.add_paragraph("CAUSE NO. 2024-CR-1042-D")
    docx_doc.add_paragraph("THE STATE OF TEXAS VS. CARLOS RAMIREZ")
    docx_doc.add_paragraph("ORDER OF DEFERRED ADJUDICATION")
    docx_doc.add_paragraph("Presiding Judge: Hon. Jack W. Pulcher. On this day came on to be considered...")

    # Add a table to test table extraction
    table = docx_doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Charge"
    table.rows[0].cells[1].text = "Level"
    table.rows[0].cells[2].text = "Disposition"
    table.rows[1].cells[0].text = "POSS CS PG 1 <1G"
    table.rows[1].cells[1].text = "State Jail Felony"
    table.rows[1].cells[2].text = "Deferred Adjudication"

    docx_bytes_io = io.BytesIO()
    docx_doc.save(docx_bytes_io)
    docx_bytes = docx_bytes_io.getvalue()

    h_docx = validate_document_health(docx_bytes, "2024-CR-1042-D_Order.docx")
    assert h_docx["is_healthy"], f"Valid DOCX should be healthy: {h_docx}"
    print(f"  [OK] DOCX health check passed ({len(docx_bytes)} bytes, {h_docx['detected_format']})")

    # Save temp and extract
    temp_docx = "temp_test_order.docx"
    with open(temp_docx, "wb") as f:
        f.write(docx_bytes)
    try:
        extracted_docx = extract_content(temp_docx, "2024-CR-1042-D_Order.docx")
        assert "105TH DISTRICT COURT" in extracted_docx["text"]
        assert len(extracted_docx["tables"]) > 0
        print(f"  [OK] Extracted {len(extracted_docx['text'])} chars and {len(extracted_docx['tables'])} table(s) from DOCX via {extracted_docx['extraction_method']}")

        legal_docx = extract_legal_entities(extracted_docx["text"], "2024-CR-1042-D_Order.docx")
        assert legal_docx["primary_case_number"] == "2024-CR-1042-D"
        assert legal_docx["court"] == "105th District Court"
        assert legal_docx["classification_label"] == "court_order"
        print(f"  [OK] Legal Entities: Case #{legal_docx['primary_case_number']}, Court: {legal_docx['court']}, Type: {legal_docx['classification_label']} (Confidence: {legal_docx['confidence']})")
    finally:
        if os.path.exists(temp_docx):
            os.remove(temp_docx)

    print("\n=== Step 3: Testing Excel (XLSX) Generation, Health & Extraction ===")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Appointed Cases Roster"
    ws.append(["Case Number", "Defendant", "Court", "Disposition Date", "Voucher Amount"])
    ws.append(["2024-CR-2180-C", "Elena Trevino", "94th District Court", "2024-05-10", 1250.0])
    ws.append(["2023-CR-4819-A", "Marcus Hernandez", "28th District Court", "2024-02-14", 2450.0])
    ws.append(["2024-CR-3102-E", "David Garcia", "214th District Court", "2024-07-20", 2100.0])

    xlsx_bytes_io = io.BytesIO()
    wb.save(xlsx_bytes_io)
    xlsx_bytes = xlsx_bytes_io.getvalue()

    h_xlsx = validate_document_health(xlsx_bytes, "Nueces_Master_Roster.xlsx")
    assert h_xlsx["is_healthy"], f"Valid XLSX should be healthy: {h_xlsx}"
    print(f"  [OK] XLSX health check passed ({len(xlsx_bytes)} bytes)")

    temp_xlsx = "temp_test_roster.xlsx"
    with open(temp_xlsx, "wb") as f:
        f.write(xlsx_bytes)
    try:
        extracted_xlsx = extract_content(temp_xlsx, "Nueces_Master_Roster.xlsx")
        assert len(extracted_xlsx["sheets"]) == 1
        assert extracted_xlsx["sheets"][0]["row_count"] == 4
        print(f"  [OK] Extracted sheet '{extracted_xlsx['sheets'][0]['sheet_name']}' with {extracted_xlsx['sheets'][0]['row_count']} rows via {extracted_xlsx['extraction_method']}")

        legal_xlsx = extract_legal_entities(extracted_xlsx["text"], "Nueces_Master_Roster.xlsx")
        assert "2024-CR-2180-C" in legal_xlsx["all_case_numbers"]
        assert legal_xlsx["classification_label"] == "spreadsheet_roster"
        print(f"  [OK] Excel legal analysis: Found {len(legal_xlsx['all_case_numbers'])} case numbers: {legal_xlsx['all_case_numbers']}")
    finally:
        if os.path.exists(temp_xlsx):
            os.remove(temp_xlsx)

    print("\n=== Step 4: Testing CSV Extraction & Sniffing ===")
    csv_content = """Case Number,Defendant,Court,Charge,Status
2024-CR-0891-B,Brandon Miller,319th District Court,Burglary of Habitation,IN_CUSTODY
2023-CR-1109-B,Sophia Benavides,148th District Court,Evading Arrest,PAID
"""
    h_csv = validate_document_health(csv_content.encode("utf-8"), "cases.csv")
    assert h_csv["is_healthy"]

    temp_csv = "temp_test.csv"
    with open(temp_csv, "w", encoding="utf-8") as f:
        f.write(csv_content)
    try:
        extracted_csv = extract_content(temp_csv, "cases.csv")
        assert len(extracted_csv["sheets"][0]["rows"]) == 3
        legal_csv = extract_legal_entities(extracted_csv["text"], "cases.csv")
        assert "2024-CR-0891-B" in legal_csv["all_case_numbers"]
        print(f"  [OK] CSV extracted: {len(legal_csv['all_case_numbers'])} cases identified.")
    finally:
        if os.path.exists(temp_csv):
            os.remove(temp_csv)

    print("\n=== Step 5: Testing Order of Appointment Phone/Email & Custody Extraction ===")
    sample_order_ocr = """
    CAUSE NO. 26FC-3771B
    THE STATE OF TEXAS VS. SHAUN DOUGLAS JACKSON
    IN THE COUNTY COURT AT LAW NO. 2 OF NUECES COUNTY, TEXAS
    ORDER OF APPOINTMENT OF COUNSEL
    The Court hereby appoints Kimbel L. Brown, Attorney at Law, SBN: 24040404
    to represent: SHAUN DOUGLAS JACKSON, DOB: 05/14/1988, SO NO: 994821
    Charge: POSS CS PG 1/1-B <1G, State Jail Felony
    currently in Jail: Yes
    Address: 4809 CALALLEN DR, CORPUS CHRISTI TX 78410
    Home Phone: 361-850-3935
    Email: SHAUNJ774@GMAIL.COM
    Signed on this the 12th day of August, 2026.
    """
    legal_order = extract_legal_entities(sample_order_ocr, "ORDER OF APPOINTMENT-26FC-3771B.pdf")
    assert legal_order["primary_case_number"] == "26FC-3771B"
    assert legal_order["defendant_name"] == "Shaun Douglas Jackson"
    assert legal_order["court"] == "County Court at Law No. 2"
    assert legal_order["specialized_fields"]["phone"] == "(361) 850-3935", f"Expected normalized phone, got {legal_order['specialized_fields'].get('phone')}"
    assert legal_order["specialized_fields"]["email"] == "shaunj774@gmail.com", f"Expected clean email without phone bleed, got {legal_order['specialized_fields'].get('email')}"
    assert legal_order["specialized_fields"]["in_custody"] is True, "Expected in_custody == True"
    assert legal_order["specialized_fields"]["dob"] == "05/14/1988"
    print(f"  [OK] Extracted Phone: {legal_order['specialized_fields']['phone']} (normalized format)")
    print(f"  [OK] Extracted Email: {legal_order['specialized_fields']['email']} (phone bleed prevented)")
    print(f"  [OK] Extracted Custody Status: in_custody={legal_order['specialized_fields']['in_custody']}")

    # Validate Pydantic Schema checks
    from app.schemas.client import ClientCreate, ClientBase
    valid_c = ClientCreate(
        name="Shaun Douglas Jackson",
        phone="3618503935",
        email="361-850-3935shaunj774@gmail.com"  # Test auto-clean on schema ingestion
    )
    assert valid_c.phone == "(361) 850-3935"
    assert valid_c.email == "shaunj774@gmail.com"
    print("  [OK] Pydantic ClientCreate normalized phone and cleaned merged email prefix")

    # Invalid phone rejection
    try:
        ClientCreate(name="Bad Phone", phone="12345")
        assert False, "Should have raised ValueError on invalid phone"
    except ValueError:
        print("  [OK] Pydantic rejected phone with invalid digit count (<10 digits)")

    # Invalid email rejection
    try:
        ClientCreate(name="Bad Email", email="not-an-email")
        assert False, "Should have raised ValueError on invalid email"
    except ValueError:
        print("  [OK] Pydantic rejected malformed email syntax")

    print("\n=== Step 6: Testing Acceptance of Appointment Extraction & Review Fallback ===")
    sample_acceptance_path = "C:/antigravity/Filestavk/Acceptance of Appointment-26MC-02456.pdf"
    if os.path.exists(sample_acceptance_path):
        extracted_acc = extract_content(sample_acceptance_path, "Acceptance of Appointment-26MC-02456.pdf")
        legal_acc = extract_legal_entities(extracted_acc["text"], "Acceptance of Appointment-26MC-02456.pdf")
        assert legal_acc["classification_label"] == "appointment_acceptance"
        assert legal_acc["primary_case_number"] == "26MC-02456"
        assert legal_acc["defendant_name"] == "Oscar Perez"
        spec = legal_acc["specialized_fields"]
        assert spec["has_clerk_file_stamp"] is True
        assert spec["has_attorney_affirmation"] is True
        assert spec["needs_human_review"] is False
        assert spec["client_contact_date"] == "09/01/2026"
        assert "September 02, 2026" in spec["acceptance_filed_date"]
        print(f"  [OK] Real PDF Ingested: Case #{legal_acc['primary_case_number']} ({legal_acc['defendant_name']})")
        print(f"  [OK] Clerk Stamp Detected: {spec['has_clerk_file_stamp']} ({spec['acceptance_filed_date']})")
        print(f"  [OK] Attorney Affirmation Detected: {spec['has_attorney_affirmation']} (Contact: {spec['client_contact_date']})")
        print(f"  [OK] Needs Human Review: {spec['needs_human_review']}")

    # Test blank acceptance template without stamp or affirmation -> must trigger needs_human_review = True
    blank_template_text = """
    CAUSE NO. 26MC-99999
    THE STATE OF TEXAS VS. JOHN DOE
    IN THE COUNTY COURT AT LAW NO. 1 OF NUECES COUNTY, TEXAS
    ORDER OF APPOINTMENT OF COUNSEL
    The Court hereby appoints Kimbel L. Brown, Attorney at Law...
    AFFIRMED:
    Attorney for Defendant
    Date: / /
    """
    legal_blank = extract_legal_entities(blank_template_text, "Acceptance of Appointment-26MC-99999.pdf")
    spec_blank = legal_blank["specialized_fields"]
    assert spec_blank["has_clerk_file_stamp"] is False
    assert spec_blank["has_attorney_affirmation"] is False
    assert spec_blank["needs_human_review"] is True
    assert "human_review_instructions" in spec_blank
    print(f"  [OK] Blank template correctly flagged for human review: needs_human_review={spec_blank['needs_human_review']}")

    print("\n>>> ALL DOCUMENT EXTRACTION & HEALTH VALIDATION TESTS PASSED! <<<")

if __name__ == "__main__":
    run_tests()

